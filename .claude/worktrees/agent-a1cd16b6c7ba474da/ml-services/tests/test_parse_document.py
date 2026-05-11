"""Tests for document parsing endpoint (W37).

Uses python-docx to generate a DOCX in memory — no fixture files needed.
"""

import io
from fastapi.testclient import TestClient
from app.main import app

client = TestClient(app)


def make_test_docx() -> bytes:
    """Generate a minimal DOCX file in memory."""
    import docx

    doc = docx.Document()
    doc.add_heading("Project Report", level=1)
    doc.add_paragraph("This is the introduction to the project report.")
    doc.add_heading("Methodology", level=2)
    doc.add_paragraph("We used a mixed-methods approach combining qualitative and quantitative analysis.")
    doc.add_heading("Results", level=2)
    doc.add_paragraph("The results show a 25% improvement in processing speed.")

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf.read()


def test_parse_docx():
    data = make_test_docx()
    resp = client.post(
        "/parse-document",
        files={"file": ("report.docx", data, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
    )
    assert resp.status_code == 200
    result = resp.json()
    assert result["format"] == "docx"
    assert result["title"] == "Project Report"
    assert len(result["sections"]) >= 2
    assert result["word_count"] > 10


def test_docx_sections_have_headings():
    data = make_test_docx()
    resp = client.post(
        "/parse-document",
        files={"file": ("report.docx", data, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
    )
    result = resp.json()
    headings = [s["heading"] for s in result["sections"] if s.get("heading")]
    assert "Project Report" in headings or "Methodology" in headings


def test_unsupported_file_type():
    resp = client.post(
        "/parse-document",
        files={"file": ("data.csv", b"a,b,c\n1,2,3", "text/csv")},
    )
    assert resp.status_code == 400
    assert "Unsupported" in resp.json()["detail"]
